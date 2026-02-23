#!/usr/bin/env python3
"""
Zero Day Fortress — Hackathon Presentation Generator
Creates a professional, visually striking PowerPoint presentation.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── Color Palette (Cyber/Dark Theme) ──────────────────────────────────────
BG_DARK = RGBColor(0x0A, 0x0A, 0x1A)       # Deep navy-black
BG_CARD = RGBColor(0x12, 0x14, 0x2A)       # Slightly lighter card bg
ACCENT_CYAN = RGBColor(0x00, 0xD4, 0xFF)   # Neon cyan
ACCENT_RED = RGBColor(0xFF, 0x3E, 0x3E)    # Attack red
ACCENT_GREEN = RGBColor(0x00, 0xFF, 0x88)  # Defense green
ACCENT_PURPLE = RGBColor(0xA855, 0xF7, 0x00)[0:3] if False else RGBColor(0xA8, 0x55, 0xF7)  # Purple
ACCENT_YELLOW = RGBColor(0xFF, 0xD6, 0x00) # Warning yellow
ACCENT_ORANGE = RGBColor(0xFF, 0x8C, 0x00) # Orange
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xB0, 0xB0, 0xC0)
DIM_GRAY = RGBColor(0x70, 0x70, 0x90)
CRITICAL_RED = RGBColor(0xFF, 0x00, 0x44)
HIGH_ORANGE = RGBColor(0xFF, 0x6B, 0x35)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_bg(slide, color=BG_DARK):
    """Fill slide background with solid color."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, fill_color=None, border_color=None, border_width=Pt(1), shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    """Add a rounded rectangle shape."""
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.shadow.inherit = False
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
    else:
        shape.line.fill.background()
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=18, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    """Add a text box with styled text."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_multi_text(slide, left, top, width, height, lines, default_size=16, default_color=LIGHT_GRAY):
    """Add text box with multiple styled paragraphs. Each line is (text, size, color, bold)."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line_data in enumerate(lines):
        text = line_data[0]
        size = line_data[1] if len(line_data) > 1 else default_size
        color = line_data[2] if len(line_data) > 2 else default_color
        bold = line_data[3] if len(line_data) > 3 else False
        align = line_data[4] if len(line_data) > 4 else PP_ALIGN.LEFT
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.name = "Calibri"
        p.alignment = align
        p.space_after = Pt(4)
    return txBox


def add_accent_line(slide, left, top, width, color=ACCENT_CYAN):
    """Add a thin accent line."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_badge(slide, left, top, text, bg_color, text_color=WHITE, width=Inches(1.5), height=Inches(0.4), font_size=11):
    """Add a colored badge/pill."""
    shape = add_shape(slide, left, top, width, height, fill_color=bg_color)
    shape.text_frame.word_wrap = True
    p = shape.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = text_color
    p.font.bold = True
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER
    shape.text_frame.paragraphs[0].space_before = Pt(0)
    shape.text_frame.paragraphs[0].space_after = Pt(0)
    return shape


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 1: TITLE SLIDE
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
add_bg(slide)

# Top accent bar
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_CYAN)

# Shield icon area (geometric shape)
add_shape(slide, Inches(5.9), Inches(1.2), Inches(1.5), Inches(1.5),
          fill_color=None, border_color=ACCENT_CYAN, border_width=Pt(3),
          shape_type=MSO_SHAPE.DIAMOND)

# Title
add_text_box(slide, Inches(1), Inches(3.0), Inches(11.3), Inches(1.2),
             "ZERO DAY FORTRESS", font_size=54, color=WHITE, bold=True,
             alignment=PP_ALIGN.CENTER, font_name="Calibri")

# Accent line under title
add_accent_line(slide, Inches(4.5), Inches(4.2), Inches(4.3), ACCENT_CYAN)

# Tagline
add_text_box(slide, Inches(1.5), Inches(4.5), Inches(10.3), Inches(0.8),
             "Autonomous AI Red vs Blue Team Cyber Defense Platform",
             font_size=24, color=ACCENT_CYAN, bold=False, alignment=PP_ALIGN.CENTER)

# Sub-tagline
add_text_box(slide, Inches(2), Inches(5.3), Inches(9.3), Inches(0.6),
             "Continuous AI-driven cyber war inside your codebase — so real hackers never win.",
             font_size=16, color=LIGHT_GRAY, bold=False, alignment=PP_ALIGN.CENTER)

# Bottom badges
add_badge(slide, Inches(3.2), Inches(6.3), "AI-POWERED", ACCENT_CYAN, BG_DARK, width=Inches(1.6))
add_badge(slide, Inches(5.0), Inches(6.3), "MULTI-LLM", ACCENT_PURPLE, WHITE, width=Inches(1.5))
add_badge(slide, Inches(6.7), Inches(6.3), "REAL-TIME", ACCENT_GREEN, BG_DARK, width=Inches(1.5))
add_badge(slide, Inches(8.4), Inches(6.3), "AUTONOMOUS", ACCENT_RED, WHITE, width=Inches(1.7))


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 2: THE PROBLEM
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_RED)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(6), Inches(0.7),
             "THE PROBLEM", font_size=36, color=ACCENT_RED, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2.5), ACCENT_RED)

# Problem cards
problems = [
    ("$4.88M", "Average cost of a\ndata breach in 2024", CRITICAL_RED),
    ("68%", "Of breaches involve\nhuman error", HIGH_ORANGE),
    ("277 Days", "Average time to identify\n& contain a breach", ACCENT_YELLOW),
    ("3,500+", "New CVEs published\nevery month", ACCENT_RED),
]

for i, (stat, desc, color) in enumerate(problems):
    x = Inches(0.6 + i * 3.1)
    y = Inches(1.6)
    card = add_shape(slide, x, y, Inches(2.8), Inches(2.8), fill_color=BG_CARD, border_color=color, border_width=Pt(2))
    add_text_box(slide, x + Inches(0.2), y + Inches(0.3), Inches(2.4), Inches(1.0),
                 stat, font_size=42, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.2), y + Inches(1.5), Inches(2.4), Inches(1.0),
                 desc, font_size=15, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# Bottom insight
add_shape(slide, Inches(0.6), Inches(4.8), Inches(12.1), Inches(2.0), fill_color=BG_CARD, border_color=DIM_GRAY)
add_multi_text(slide, Inches(1.0), Inches(5.0), Inches(11.3), Inches(1.6), [
    ("Current Security Tools Fall Short:", 22, ACCENT_YELLOW, True),
    ("", 8, DIM_GRAY),
    ("•  Manual penetration testing is expensive, slow, and infrequent", 17, LIGHT_GRAY),
    ("•  Static analyzers produce too many false positives — alert fatigue kills response time", 17, LIGHT_GRAY),
    ("•  Most tools find vulnerabilities but DON'T fix them — developers still carry the burden", 17, LIGHT_GRAY),
    ("•  No tool provides continuous, autonomous attack + defense + compliance in one workflow", 17, LIGHT_GRAY),
])


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 3: OUR SOLUTION
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_GREEN)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
             "OUR SOLUTION: ZERO DAY FORTRESS", font_size=36, color=ACCENT_GREEN, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(3.5), ACCENT_GREEN)

add_text_box(slide, Inches(0.8), Inches(1.4), Inches(11.5), Inches(0.8),
             "An autonomous AI platform that attacks your code, fixes every vulnerability, and proves the fixes work — all in one continuous battle.",
             font_size=18, color=LIGHT_GRAY)

# Three pillars
pillars = [
    ("ATTACK", "AI Red Team agents\nautonomously scan, exploit,\nand chain vulnerabilities\nin your codebase.", ACCENT_RED, "Spider • Blade\nVenom • Phantom"),
    ("DEFEND", "AI Blue Team agents\ngenerate patches, verify fixes,\nharden infrastructure,\nand ensure compliance.", ACCENT_GREEN, "Shield • Proof\nFortress • Auditor"),
    ("COMPARE", "Battle Royale pits multiple\nLLMs against each other.\nCross-Examination Arena lets\nAIs debate findings.", ACCENT_PURPLE, "Claude • GPT-4\nGemini • Ollama"),
]

for i, (title, desc, color, agents) in enumerate(pillars):
    x = Inches(0.6 + i * 4.2)
    y = Inches(2.5)
    card = add_shape(slide, x, y, Inches(3.8), Inches(4.5), fill_color=BG_CARD, border_color=color, border_width=Pt(2))

    # Title
    add_text_box(slide, x + Inches(0.3), y + Inches(0.3), Inches(3.2), Inches(0.6),
                 title, font_size=28, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_accent_line(slide, x + Inches(0.8), y + Inches(1.0), Inches(2.2), color)

    # Description
    add_text_box(slide, x + Inches(0.3), y + Inches(1.2), Inches(3.2), Inches(1.8),
                 desc, font_size=15, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    # Agent names
    add_shape(slide, x + Inches(0.5), y + Inches(3.2), Inches(2.8), Inches(1.0),
              fill_color=BG_DARK, border_color=color, border_width=Pt(1))
    add_text_box(slide, x + Inches(0.5), y + Inches(3.3), Inches(2.8), Inches(0.9),
                 agents, font_size=13, color=color, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 4: HOW IT WORKS — BATTLE LIFECYCLE
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_CYAN)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
             "HOW IT WORKS — BATTLE LIFECYCLE", font_size=36, color=ACCENT_CYAN, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(3.5), ACCENT_CYAN)

phases = [
    ("1", "RECON", "Spider maps attack\nsurface, endpoints,\nauth flows, secrets", ACCENT_CYAN),
    ("2", "ASSAULT", "Blade exploits vulns\nPhantom finds logic\nflaws (parallel)", ACCENT_RED),
    ("3", "CHAINING", "Venom chains low-risk\nvulns into critical\nmulti-step attacks", ACCENT_ORANGE),
    ("4", "PATCHING", "Shield generates\nminimal secure patches\nProof verifies them", ACCENT_GREEN),
    ("5", "HARDENING", "Fortress applies\nsystemic defenses:\nheaders, CORS, etc.", ACCENT_PURPLE),
    ("6", "REPORTING", "Auditor maps to\nOWASP, CWE, SOC2,\nPCI-DSS, GDPR", ACCENT_YELLOW),
    ("7", "COMPLETE", "Full battle summary\nwith metrics, scores,\nand PR creation", WHITE),
]

for i, (num, name, desc, color) in enumerate(phases):
    x = Inches(0.3 + i * 1.85)
    y = Inches(1.6)

    # Phase card
    card = add_shape(slide, x, y, Inches(1.7), Inches(3.6), fill_color=BG_CARD, border_color=color, border_width=Pt(2))

    # Phase number circle
    circle = add_shape(slide, x + Inches(0.55), y + Inches(0.2), Inches(0.6), Inches(0.6),
                       fill_color=color, shape_type=MSO_SHAPE.OVAL)
    circle.text_frame.paragraphs[0].text = num
    circle.text_frame.paragraphs[0].font.size = Pt(22)
    circle.text_frame.paragraphs[0].font.color.rgb = BG_DARK
    circle.text_frame.paragraphs[0].font.bold = True
    circle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Phase name
    add_text_box(slide, x + Inches(0.1), y + Inches(1.0), Inches(1.5), Inches(0.5),
                 name, font_size=14, color=color, bold=True, alignment=PP_ALIGN.CENTER)

    # Phase description
    add_text_box(slide, x + Inches(0.1), y + Inches(1.5), Inches(1.5), Inches(1.8),
                 desc, font_size=12, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    # Arrow between phases
    if i < 6:
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + Inches(1.72), y + Inches(1.5), Inches(0.15), Inches(0.3))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = DIM_GRAY
        arrow.line.fill.background()
        arrow.shadow.inherit = False

# Bottom note
add_text_box(slide, Inches(0.8), Inches(5.6), Inches(11.5), Inches(1.2),
             "The entire lifecycle runs autonomously — from first scan to verified patches — in a single command.\nReal-time WebSocket updates stream every agent action to the live dashboard.",
             font_size=15, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 5: RED TEAM — THE ATTACKERS
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_RED)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
             "RED TEAM — THE ATTACKERS", font_size=36, color=ACCENT_RED, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(3.0), ACCENT_RED)

red_agents = [
    ("SPIDER", "Recon Agent", [
        "Maps all endpoints & input vectors",
        "Identifies auth flows & session mgmt",
        "Finds database queries & ORM patterns",
        "Detects secret exposures in code",
        "Catalogs external API calls",
        "Identifies file upload/download handlers",
    ]),
    ("BLADE", "Exploiter Agent", [
        "Generates exploit payloads per vuln type",
        "Tests SQL Injection (auth bypass, union)",
        "Crafts XSS payloads (reflected, stored)",
        "Probes SSRF, IDOR, Path Traversal",
        "Attempts Command Injection attacks",
        "Validates with proof-of-concept code",
    ]),
    ("VENOM", "Chain Builder", [
        "Chains low-severity → critical attacks",
        "Maps multi-step attack narratives",
        "Calculates combined risk scores",
        "Example: Info leak → IDOR → Takeover",
        "Identifies shared vulnerable components",
        "Produces exploit chain visualizations",
    ]),
    ("PHANTOM", "Business Logic", [
        "Finds price manipulation flaws",
        "Detects coupon/discount abuse paths",
        "Identifies workflow bypass scenarios",
        "Discovers race condition windows",
        "Checks privilege escalation paths",
        "Tests negative quantity / boundary cases",
    ]),
]

for i, (name, role, capabilities) in enumerate(red_agents):
    x = Inches(0.4 + i * 3.2)
    y = Inches(1.5)
    card = add_shape(slide, x, y, Inches(3.0), Inches(5.3), fill_color=BG_CARD, border_color=ACCENT_RED, border_width=Pt(2))

    # Agent name
    add_text_box(slide, x + Inches(0.2), y + Inches(0.2), Inches(2.6), Inches(0.5),
                 name, font_size=24, color=ACCENT_RED, bold=True, alignment=PP_ALIGN.CENTER)
    # Role
    add_text_box(slide, x + Inches(0.2), y + Inches(0.7), Inches(2.6), Inches(0.4),
                 role, font_size=14, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)
    add_accent_line(slide, x + Inches(0.4), y + Inches(1.15), Inches(2.2), ACCENT_RED)

    # Capabilities
    lines = [(cap, 12, LIGHT_GRAY) for cap in capabilities]
    add_multi_text(slide, x + Inches(0.3), y + Inches(1.3), Inches(2.5), Inches(3.8), lines, default_size=12)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 6: BLUE TEAM — THE DEFENDERS
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_GREEN)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
             "BLUE TEAM — THE DEFENDERS", font_size=36, color=ACCENT_GREEN, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(3.0), ACCENT_GREEN)

blue_agents = [
    ("SHIELD", "Patcher Agent", [
        "Generates minimal secure patches",
        "Language-aware code transformations",
        "One-line explanations for each fix",
        "Parameterized queries for SQLi",
        "Input sanitization for XSS",
        "Access control fixes for IDOR",
    ]),
    ("PROOF", "Verification Agent", [
        "Re-runs exploits against patches",
        "Confirms vulnerability is resolved",
        "Tests for regression issues",
        "Validates patch completeness",
        "Ensures no new vulns introduced",
        "Provides verification evidence",
    ]),
    ("FORTRESS", "Hardening Agent", [
        "Adds rate limiting configurations",
        "Sets security headers (CSP, HSTS)",
        "Configures strict CORS policies",
        "Enables security audit logging",
        "Implements secrets management",
        "Adds input validation middleware",
    ]),
    ("AUDITOR", "Compliance Agent", [
        "Maps to OWASP Top 10 (2021)",
        "Assigns CWE identifiers",
        "Checks SOC2 Trust Criteria",
        "Validates PCI-DSS requirements",
        "Reviews GDPR data protection",
        "Before/after compliance scoring",
    ]),
]

for i, (name, role, capabilities) in enumerate(blue_agents):
    x = Inches(0.4 + i * 3.2)
    y = Inches(1.5)
    card = add_shape(slide, x, y, Inches(3.0), Inches(5.3), fill_color=BG_CARD, border_color=ACCENT_GREEN, border_width=Pt(2))

    add_text_box(slide, x + Inches(0.2), y + Inches(0.2), Inches(2.6), Inches(0.5),
                 name, font_size=24, color=ACCENT_GREEN, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.2), y + Inches(0.7), Inches(2.6), Inches(0.4),
                 role, font_size=14, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)
    add_accent_line(slide, x + Inches(0.4), y + Inches(1.15), Inches(2.2), ACCENT_GREEN)

    lines = [(cap, 12, LIGHT_GRAY) for cap in capabilities]
    add_multi_text(slide, x + Inches(0.3), y + Inches(1.3), Inches(2.5), Inches(3.8), lines, default_size=12)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 7: MULTI-LLM BATTLE ROYALE
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_PURPLE)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "BATTLE ROYALE — MULTI-LLM SHOWDOWN", font_size=36, color=ACCENT_PURPLE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(4.0), ACCENT_PURPLE)

add_text_box(slide, Inches(0.8), Inches(1.4), Inches(11.5), Inches(0.6),
             "Run the same battle across 4 LLM providers simultaneously. Compare who finds what. Then watch them debate.",
             font_size=17, color=LIGHT_GRAY)

# LLM Provider cards
providers = [
    ("CLAUDE", "Anthropic", "claude-sonnet-4-6", ACCENT_ORANGE),
    ("GPT-4o", "OpenAI", "gpt-4o", ACCENT_GREEN),
    ("GEMINI", "Google", "gemini-2.0-flash", ACCENT_CYAN),
    ("OLLAMA", "Local LLM", "llama3.1 / custom", ACCENT_PURPLE),
]

for i, (name, company, model, color) in enumerate(providers):
    x = Inches(0.4 + i * 3.2)
    y = Inches(2.2)
    card = add_shape(slide, x, y, Inches(3.0), Inches(2.4), fill_color=BG_CARD, border_color=color, border_width=Pt(2))

    add_text_box(slide, x + Inches(0.2), y + Inches(0.2), Inches(2.6), Inches(0.5),
                 name, font_size=26, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.2), y + Inches(0.8), Inches(2.6), Inches(0.4),
                 company, font_size=16, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.2), y + Inches(1.3), Inches(2.6), Inches(0.4),
                 model, font_size=12, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)

    # Simulated stats
    add_badge(slide, x + Inches(0.3), y + Inches(1.8), "READY", color, BG_DARK, width=Inches(2.4), height=Inches(0.35))

# Cross-Examination section
add_shape(slide, Inches(0.4), Inches(5.0), Inches(12.5), Inches(2.1), fill_color=BG_CARD, border_color=ACCENT_YELLOW, border_width=Pt(2))
add_text_box(slide, Inches(0.8), Inches(5.15), Inches(5), Inches(0.5),
             "CROSS-EXAMINATION ARENA", font_size=22, color=ACCENT_YELLOW, bold=True)

add_multi_text(slide, Inches(0.8), Inches(5.7), Inches(5.5), Inches(1.3), [
    ("Two LLMs debate each other's findings live:", 15, LIGHT_GRAY, True),
    ("", 6, DIM_GRAY),
    ("1. Challenger questions severity & patch quality", 14, LIGHT_GRAY),
    ("2. Defender argues or concedes with evidence", 14, LIGHT_GRAY),
    ("3. Verdict determines final classification", 14, LIGHT_GRAY),
])

# Debate visualization
add_shape(slide, Inches(7.0), Inches(5.3), Inches(2.5), Inches(1.5), fill_color=BG_DARK, border_color=ACCENT_RED, border_width=Pt(1))
add_text_box(slide, Inches(7.0), Inches(5.4), Inches(2.5), Inches(0.4),
             "CHALLENGER", font_size=14, color=ACCENT_RED, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(7.0), Inches(5.8), Inches(2.5), Inches(0.8),
             '"Is this really critical?\nThe CVSS score seems\ninflated..."', font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(9.6), Inches(5.9), Inches(0.5), Inches(0.5),
             "VS", font_size=18, color=ACCENT_YELLOW, bold=True, alignment=PP_ALIGN.CENTER)

add_shape(slide, Inches(10.2), Inches(5.3), Inches(2.5), Inches(1.5), fill_color=BG_DARK, border_color=ACCENT_GREEN, border_width=Pt(1))
add_text_box(slide, Inches(10.2), Inches(5.4), Inches(2.5), Inches(0.4),
             "DEFENDER", font_size=14, color=ACCENT_GREEN, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(10.2), Inches(5.8), Inches(2.5), Inches(0.8),
             '"The chain from info\nleak to account takeover\nconfirms critical."', font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 8: TECH STACK
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_CYAN)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(6), Inches(0.7),
             "TECH STACK", font_size=36, color=ACCENT_CYAN, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2.0), ACCENT_CYAN)

# Backend column
add_shape(slide, Inches(0.4), Inches(1.5), Inches(4.0), Inches(5.5), fill_color=BG_CARD, border_color=ACCENT_CYAN, border_width=Pt(2))
add_text_box(slide, Inches(0.6), Inches(1.6), Inches(3.6), Inches(0.5),
             "BACKEND", font_size=22, color=ACCENT_CYAN, bold=True, alignment=PP_ALIGN.CENTER)
add_accent_line(slide, Inches(1.2), Inches(2.15), Inches(2.4), ACCENT_CYAN)

backend_items = [
    ("Python 3.12", "Core language", ACCENT_CYAN),
    ("FastAPI", "REST & WebSocket API", ACCENT_CYAN),
    ("Anthropic SDK", "Claude integration", ACCENT_ORANGE),
    ("OpenAI SDK", "GPT-4 integration", ACCENT_GREEN),
    ("Google GenAI", "Gemini integration", ACCENT_CYAN),
    ("Bandit + Semgrep", "Security analysis", ACCENT_RED),
    ("GitPython", "Repository operations", LIGHT_GRAY),
    ("Pydantic", "Data validation", LIGHT_GRAY),
]
for j, (tech, desc, color) in enumerate(backend_items):
    yy = Inches(2.3 + j * 0.55)
    add_text_box(slide, Inches(0.8), yy, Inches(1.8), Inches(0.4), tech, font_size=13, color=color, bold=True)
    add_text_box(slide, Inches(2.6), yy, Inches(1.8), Inches(0.4), desc, font_size=12, color=DIM_GRAY)

# Frontend column
add_shape(slide, Inches(4.7), Inches(1.5), Inches(4.0), Inches(5.5), fill_color=BG_CARD, border_color=ACCENT_GREEN, border_width=Pt(2))
add_text_box(slide, Inches(4.9), Inches(1.6), Inches(3.6), Inches(0.5),
             "FRONTEND", font_size=22, color=ACCENT_GREEN, bold=True, alignment=PP_ALIGN.CENTER)
add_accent_line(slide, Inches(5.5), Inches(2.15), Inches(2.4), ACCENT_GREEN)

frontend_items = [
    ("Next.js 14", "React framework", ACCENT_GREEN),
    ("TypeScript", "Type safety", ACCENT_CYAN),
    ("Tailwind CSS", "Utility-first styling", ACCENT_PURPLE),
    ("Recharts", "Data visualization", ACCENT_ORANGE),
    ("Force Graph", "Network visualization", ACCENT_YELLOW),
    ("WebSocket", "Real-time streaming", ACCENT_GREEN),
    ("Lucide Icons", "438+ SVG icons", LIGHT_GRAY),
    ("React 18", "UI library", LIGHT_GRAY),
]
for j, (tech, desc, color) in enumerate(frontend_items):
    yy = Inches(2.3 + j * 0.55)
    add_text_box(slide, Inches(5.1), yy, Inches(1.8), Inches(0.4), tech, font_size=13, color=color, bold=True)
    add_text_box(slide, Inches(6.9), yy, Inches(1.8), Inches(0.4), desc, font_size=12, color=DIM_GRAY)

# Infrastructure column
add_shape(slide, Inches(9.0), Inches(1.5), Inches(4.0), Inches(5.5), fill_color=BG_CARD, border_color=ACCENT_PURPLE, border_width=Pt(2))
add_text_box(slide, Inches(9.2), Inches(1.6), Inches(3.6), Inches(0.5),
             "INFRASTRUCTURE", font_size=22, color=ACCENT_PURPLE, bold=True, alignment=PP_ALIGN.CENTER)
add_accent_line(slide, Inches(9.8), Inches(2.15), Inches(2.4), ACCENT_PURPLE)

infra_items = [
    ("Docker", "Containerization", ACCENT_PURPLE),
    ("Docker Compose", "Multi-service orchestration", ACCENT_PURPLE),
    ("Uvicorn", "ASGI server", ACCENT_CYAN),
    ("GitHub API", "PR automation", LIGHT_GRAY),
    ("Makefile", "Build automation", LIGHT_GRAY),
    ("3 Services", "Backend + Frontend + Target", ACCENT_YELLOW),
    ("Port 8000", "Backend API", DIM_GRAY),
    ("Port 3000", "Frontend UI", DIM_GRAY),
]
for j, (tech, desc, color) in enumerate(infra_items):
    yy = Inches(2.3 + j * 0.55)
    add_text_box(slide, Inches(9.4), yy, Inches(1.8), Inches(0.4), tech, font_size=13, color=color, bold=True)
    add_text_box(slide, Inches(11.2), yy, Inches(1.8), Inches(0.4), desc, font_size=12, color=DIM_GRAY)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 9: ARCHITECTURE DIAGRAM
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_CYAN)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
             "SYSTEM ARCHITECTURE", font_size=36, color=ACCENT_CYAN, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(3.0), ACCENT_CYAN)

# User/Browser
add_shape(slide, Inches(5.4), Inches(1.4), Inches(2.5), Inches(0.8), fill_color=BG_CARD, border_color=ACCENT_CYAN, border_width=Pt(2))
add_text_box(slide, Inches(5.4), Inches(1.45), Inches(2.5), Inches(0.7),
             "User / Browser", font_size=16, color=ACCENT_CYAN, bold=True, alignment=PP_ALIGN.CENTER)

# Arrow down
add_shape(slide, Inches(6.5), Inches(2.2), Inches(0.3), Inches(0.4), fill_color=DIM_GRAY, shape_type=MSO_SHAPE.DOWN_ARROW)

# Frontend box
add_shape(slide, Inches(4.0), Inches(2.7), Inches(5.3), Inches(1.2), fill_color=BG_CARD, border_color=ACCENT_GREEN, border_width=Pt(2))
add_text_box(slide, Inches(4.2), Inches(2.75), Inches(4.9), Inches(0.4),
             "FRONTEND — Next.js 14 + TypeScript", font_size=16, color=ACCENT_GREEN, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(4.2), Inches(3.2), Inches(4.9), Inches(0.5),
             "Dashboard • Battle Map • Agent Terminal • Vuln Cards • Charts", font_size=12, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)

# Arrow down
add_shape(slide, Inches(6.5), Inches(3.95), Inches(0.3), Inches(0.4), fill_color=DIM_GRAY, shape_type=MSO_SHAPE.DOWN_ARROW)

# Backend box (main)
add_shape(slide, Inches(1.0), Inches(4.4), Inches(11.3), Inches(2.7), fill_color=BG_CARD, border_color=ACCENT_CYAN, border_width=Pt(2))
add_text_box(slide, Inches(1.2), Inches(4.45), Inches(10.9), Inches(0.4),
             "BACKEND — FastAPI + Python 3.12", font_size=16, color=ACCENT_CYAN, bold=True, alignment=PP_ALIGN.CENTER)

# Sub-boxes inside backend
# Orchestrator
add_shape(slide, Inches(1.3), Inches(5.0), Inches(2.2), Inches(0.7), fill_color=BG_DARK, border_color=ACCENT_YELLOW)
add_text_box(slide, Inches(1.3), Inches(5.05), Inches(2.2), Inches(0.6),
             "War General\n(Orchestrator)", font_size=11, color=ACCENT_YELLOW, alignment=PP_ALIGN.CENTER)

# Red team box
add_shape(slide, Inches(3.8), Inches(5.0), Inches(3.5), Inches(1.8), fill_color=BG_DARK, border_color=ACCENT_RED)
add_text_box(slide, Inches(3.8), Inches(5.0), Inches(3.5), Inches(0.35),
             "RED TEAM", font_size=12, color=ACCENT_RED, bold=True, alignment=PP_ALIGN.CENTER)
red_names = ["Spider", "Blade", "Venom", "Phantom"]
for k, name in enumerate(red_names):
    xx = Inches(3.9 + (k % 2) * 1.7)
    yy = Inches(5.4 + (k // 2) * 0.65)
    add_badge(slide, xx, yy, name, ACCENT_RED, WHITE, width=Inches(1.5), height=Inches(0.35), font_size=10)

# Blue team box
add_shape(slide, Inches(7.6), Inches(5.0), Inches(3.5), Inches(1.8), fill_color=BG_DARK, border_color=ACCENT_GREEN)
add_text_box(slide, Inches(7.6), Inches(5.0), Inches(3.5), Inches(0.35),
             "BLUE TEAM", font_size=12, color=ACCENT_GREEN, bold=True, alignment=PP_ALIGN.CENTER)
blue_names = ["Shield", "Proof", "Fortress", "Auditor"]
for k, name in enumerate(blue_names):
    xx = Inches(7.7 + (k % 2) * 1.7)
    yy = Inches(5.4 + (k // 2) * 0.65)
    add_badge(slide, xx, yy, name, ACCENT_GREEN, BG_DARK, width=Inches(1.5), height=Inches(0.35), font_size=10)

# LLM Providers row at bottom of backend box
add_shape(slide, Inches(1.3), Inches(5.9), Inches(2.2), Inches(0.7), fill_color=BG_DARK, border_color=ACCENT_PURPLE)
add_text_box(slide, Inches(1.3), Inches(5.95), Inches(2.2), Inches(0.6),
             "LLM Layer\nClaude • GPT • Gemini • Ollama", font_size=10, color=ACCENT_PURPLE, alignment=PP_ALIGN.CENTER)

# Target App (side)
add_shape(slide, Inches(0.4), Inches(2.7), Inches(3.2), Inches(1.2), fill_color=BG_CARD, border_color=ACCENT_ORANGE, border_width=Pt(2))
add_text_box(slide, Inches(0.4), Inches(2.75), Inches(3.2), Inches(0.4),
             "TARGET APPLICATION", font_size=14, color=ACCENT_ORANGE, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(0.4), Inches(3.2), Inches(3.2), Inches(0.5),
             "GitHub Repo / Vulnerable App\nExpress + SQLite (port 3001)", font_size=11, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)

# GitHub API (side)
add_shape(slide, Inches(9.9), Inches(2.7), Inches(3.2), Inches(1.2), fill_color=BG_CARD, border_color=LIGHT_GRAY, border_width=Pt(2))
add_text_box(slide, Inches(9.9), Inches(2.75), Inches(3.2), Inches(0.4),
             "GITHUB API", font_size=14, color=LIGHT_GRAY, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, Inches(9.9), Inches(3.2), Inches(3.2), Inches(0.5),
             "Repo Cloning\nPR Creation with Patches", font_size=11, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 10: VULNERABILITY DETECTION CAPABILITIES
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_RED)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "VULNERABILITY DETECTION CAPABILITIES", font_size=36, color=ACCENT_RED, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(4.0), ACCENT_RED)

vuln_types = [
    ("SQL Injection", "CRITICAL", "Auth bypass, data extraction,\nunion-based, blind SQLi", CRITICAL_RED),
    ("XSS", "HIGH", "Reflected, stored, DOM-based\ncross-site scripting", HIGH_ORANGE),
    ("SSRF", "HIGH", "Server-side request forgery\nto internal services", HIGH_ORANGE),
    ("IDOR", "HIGH", "Insecure direct object\nreference / broken access", HIGH_ORANGE),
    ("Path Traversal", "HIGH", "Directory traversal to read\narbitrary system files", HIGH_ORANGE),
    ("Command Injection", "CRITICAL", "OS command injection via\nunsanitized user input", CRITICAL_RED),
    ("Auth Bypass", "CRITICAL", "Authentication & session\nmanagement flaws", CRITICAL_RED),
    ("Business Logic", "MEDIUM-HIGH", "Price manipulation, coupon\nabuse, race conditions", ACCENT_YELLOW),
    ("Missing Rate Limit", "MEDIUM", "Brute force, DoS via\nunthrottled endpoints", ACCENT_YELLOW),
    ("Info Leak", "MEDIUM", "Sensitive data exposure,\nverbose errors, headers", ACCENT_YELLOW),
    ("Insecure Deps", "VARIES", "Known CVEs in third-party\npackages & libraries", DIM_GRAY),
    ("+ Attack Chains", "CRITICAL", "Multi-step attacks chaining\nlow-risk into critical", ACCENT_PURPLE),
]

for i, (name, severity, desc, color) in enumerate(vuln_types):
    col = i % 4
    row = i // 4
    x = Inches(0.4 + col * 3.2)
    y = Inches(1.5 + row * 2.0)

    card = add_shape(slide, x, y, Inches(3.0), Inches(1.7), fill_color=BG_CARD, border_color=color, border_width=Pt(1))
    add_text_box(slide, x + Inches(0.2), y + Inches(0.1), Inches(2.0), Inches(0.35),
                 name, font_size=14, color=WHITE, bold=True)
    add_badge(slide, x + Inches(1.8), y + Inches(0.12), severity, color, WHITE if color != ACCENT_YELLOW else BG_DARK,
              width=Inches(1.0), height=Inches(0.28), font_size=8)
    add_text_box(slide, x + Inches(0.2), y + Inches(0.55), Inches(2.6), Inches(1.0),
                 desc, font_size=11, color=LIGHT_GRAY)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 11: COMPLIANCE & REPORTING
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_YELLOW)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
             "COMPLIANCE & REPORTING", font_size=36, color=ACCENT_YELLOW, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(3.0), ACCENT_YELLOW)

frameworks = [
    ("OWASP\nTOP 10", "Industry standard for\nweb app security risks.\nMaps every finding to\nOWASP 2021 categories.", ACCENT_RED, "A01-A10"),
    ("CWE", "Common Weakness\nEnumeration identifiers.\nPrecise vulnerability\nclassification system.", ACCENT_ORANGE, "600+ IDs"),
    ("SOC 2", "Trust Service Criteria\nfor service organizations.\nSecurity, availability,\nprocessing integrity.", ACCENT_CYAN, "5 Categories"),
    ("PCI-DSS", "Payment Card Industry\nData Security Standard.\nRequired for handling\ncredit card data.", ACCENT_GREEN, "12 Reqs"),
    ("GDPR", "General Data Protection\nRegulation. Data exposure\nand privacy violation\nmapping.", ACCENT_PURPLE, "99 Articles"),
]

for i, (name, desc, color, stat) in enumerate(frameworks):
    x = Inches(0.3 + i * 2.6)
    y = Inches(1.5)
    card = add_shape(slide, x, y, Inches(2.4), Inches(3.5), fill_color=BG_CARD, border_color=color, border_width=Pt(2))

    add_text_box(slide, x + Inches(0.1), y + Inches(0.2), Inches(2.2), Inches(0.7),
                 name, font_size=20, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_accent_line(slide, x + Inches(0.3), y + Inches(1.0), Inches(1.8), color)
    add_text_box(slide, x + Inches(0.1), y + Inches(1.2), Inches(2.2), Inches(1.6),
                 desc, font_size=12, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
    add_badge(slide, x + Inches(0.3), y + Inches(2.9), stat, color, WHITE if color != ACCENT_YELLOW else BG_DARK,
              width=Inches(1.8), height=Inches(0.35), font_size=11)

# Compliance score visualization
add_shape(slide, Inches(0.3), Inches(5.3), Inches(12.7), Inches(1.8), fill_color=BG_CARD, border_color=ACCENT_YELLOW, border_width=Pt(1))
add_text_box(slide, Inches(0.6), Inches(5.4), Inches(4), Inches(0.5),
             "COMPLIANCE SCORE IMPROVEMENT", font_size=18, color=ACCENT_YELLOW, bold=True)

# Before bar
add_shape(slide, Inches(5.5), Inches(5.5), Inches(3.0), Inches(0.5), fill_color=ACCENT_RED, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text_box(slide, Inches(5.5), Inches(5.5), Inches(3.0), Inches(0.5),
             "BEFORE:  35/100", font_size=14, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# Arrow
add_shape(slide, Inches(8.7), Inches(5.55), Inches(0.5), Inches(0.35), fill_color=ACCENT_YELLOW, shape_type=MSO_SHAPE.RIGHT_ARROW)

# After bar
add_shape(slide, Inches(9.4), Inches(5.5), Inches(3.5), Inches(0.5), fill_color=ACCENT_GREEN, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text_box(slide, Inches(9.4), Inches(5.5), Inches(3.5), Inches(0.5),
             "AFTER:  92/100", font_size=14, color=BG_DARK, bold=True, alignment=PP_ALIGN.CENTER)

add_multi_text(slide, Inches(0.6), Inches(6.1), Inches(12.0), Inches(0.9), [
    ("Before/after scoring for every compliance framework  •  Risk reduction percentage  •  Per-finding compliance mapping", 13, DIM_GRAY, False, PP_ALIGN.LEFT),
    ("Exportable reports ready for auditors and stakeholders", 13, ACCENT_YELLOW, True, PP_ALIGN.LEFT),
])


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 12: LIVE DASHBOARD & UX
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_GREEN)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
             "LIVE DASHBOARD EXPERIENCE", font_size=36, color=ACCENT_GREEN, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(3.0), ACCENT_GREEN)

# Dashboard features
features = [
    ("Real-Time Battle Map", "WebSocket-driven live updates as\nagents discover and patch\nvulnerabilities in real-time", ACCENT_CYAN),
    ("Agent Terminal", "Watch AI agent thoughts and\nactions stream like a hacker's\nterminal — live reasoning", ACCENT_GREEN),
    ("Vulnerability Cards", "Each vuln displayed with severity\nbadge, CVSS score, affected code,\nand exploitation evidence", ACCENT_RED),
    ("Attack Chain Graph", "Interactive force-directed graph\nshowing how vulnerabilities\nchain into critical attacks", ACCENT_PURPLE),
    ("Patch Diff Viewer", "Side-by-side code diff showing\nexactly what changed in each\nsecurity patch", ACCENT_ORANGE),
    ("Phase Transitions", "Full-screen cinematic overlays\nwhen battle transitions between\nphases — immersive UX", ACCENT_YELLOW),
]

for i, (title, desc, color) in enumerate(features):
    col = i % 3
    row = i // 3
    x = Inches(0.4 + col * 4.2)
    y = Inches(1.5 + row * 2.8)

    card = add_shape(slide, x, y, Inches(3.9), Inches(2.4), fill_color=BG_CARD, border_color=color, border_width=Pt(2))
    add_text_box(slide, x + Inches(0.3), y + Inches(0.2), Inches(3.3), Inches(0.5),
                 title, font_size=18, color=color, bold=True)
    add_accent_line(slide, x + Inches(0.3), y + Inches(0.7), Inches(2.0), color)
    add_text_box(slide, x + Inches(0.3), y + Inches(0.9), Inches(3.3), Inches(1.3),
                 desc, font_size=14, color=LIGHT_GRAY)

# Bottom note
add_text_box(slide, Inches(0.8), Inches(6.8), Inches(11.5), Inches(0.5),
             "+ Boot Sequence Animation  •  Matrix Rain Background  •  Sound Effects  •  Session History  •  Provider Selector",
             font_size=13, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 13: DEMO FLOW
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_ORANGE)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(6), Inches(0.7),
             "LIVE DEMO FLOW", font_size=36, color=ACCENT_ORANGE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2.5), ACCENT_ORANGE)

demo_steps = [
    ("1", "Paste a GitHub repo URL or use\nthe built-in vulnerable app", ACCENT_CYAN, "INPUT"),
    ("2", "Select LLM providers and\nlaunch Battle Royale", ACCENT_PURPLE, "CONFIGURE"),
    ("3", "Watch Spider scan the codebase\nand map attack surface", ACCENT_CYAN, "RECON"),
    ("4", "See Blade exploit SQL injection\nwith proof-of-concept", ACCENT_RED, "EXPLOIT"),
    ("5", "Watch Shield generate patches\nand Proof verify them", ACCENT_GREEN, "PATCH"),
    ("6", "Compare findings across LLMs\nin side-by-side racing view", ACCENT_PURPLE, "COMPARE"),
    ("7", "Watch Cross-Examination Arena\nwhere LLMs debate findings", ACCENT_YELLOW, "DEBATE"),
    ("8", "View compliance report with\nOWASP/SOC2/PCI-DSS scores", ACCENT_ORANGE, "REPORT"),
]

for i, (num, desc, color, label) in enumerate(demo_steps):
    col = i % 4
    row = i // 4
    x = Inches(0.4 + col * 3.2)
    y = Inches(1.5 + row * 2.8)

    card = add_shape(slide, x, y, Inches(3.0), Inches(2.4), fill_color=BG_CARD, border_color=color, border_width=Pt(2))

    # Number circle
    circle = add_shape(slide, x + Inches(0.15), y + Inches(0.15), Inches(0.5), Inches(0.5),
                       fill_color=color, shape_type=MSO_SHAPE.OVAL)
    circle.text_frame.paragraphs[0].text = num
    circle.text_frame.paragraphs[0].font.size = Pt(18)
    circle.text_frame.paragraphs[0].font.color.rgb = BG_DARK
    circle.text_frame.paragraphs[0].font.bold = True
    circle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Label badge
    add_badge(slide, x + Inches(0.8), y + Inches(0.2), label, color, BG_DARK if color in [ACCENT_CYAN, ACCENT_GREEN, ACCENT_YELLOW] else WHITE,
              width=Inches(1.5), height=Inches(0.35), font_size=10)

    # Description
    add_text_box(slide, x + Inches(0.2), y + Inches(0.9), Inches(2.6), Inches(1.3),
                 desc, font_size=14, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 14: KEY DIFFERENTIATORS
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_CYAN)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "WHY ZERO DAY FORTRESS STANDS OUT", font_size=36, color=ACCENT_CYAN, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(4.0), ACCENT_CYAN)

differentiators = [
    ("Fully Autonomous", "From first scan to verified\npatches — zero human\nintervention required", ACCENT_GREEN),
    ("Multi-LLM Native", "First platform to pit Claude,\nGPT-4, Gemini & Ollama\nagainst each other", ACCENT_PURPLE),
    ("AI Debates AI", "Cross-Examination Arena:\nLLMs challenge each other's\nsecurity analysis live", ACCENT_YELLOW),
    ("Find + Fix + Verify", "Not just detection — auto\npatching with verification.\nComplete remediation loop.", ACCENT_CYAN),
    ("Enterprise Compliance", "Instant OWASP, SOC2,\nPCI-DSS, GDPR mapping\nwith before/after scoring", ACCENT_ORANGE),
    ("Real-Time Cinema", "Immersive dashboard with\nlive battles, animations,\nsound effects & more", ACCENT_RED),
]

for i, (title, desc, color) in enumerate(differentiators):
    col = i % 3
    row = i // 3
    x = Inches(0.4 + col * 4.2)
    y = Inches(1.5 + row * 2.8)

    card = add_shape(slide, x, y, Inches(3.9), Inches(2.4), fill_color=BG_CARD, border_color=color, border_width=Pt(2))
    add_text_box(slide, x + Inches(0.3), y + Inches(0.3), Inches(3.3), Inches(0.5),
                 title, font_size=22, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_accent_line(slide, x + Inches(0.8), y + Inches(0.85), Inches(2.3), color)
    add_text_box(slide, x + Inches(0.3), y + Inches(1.0), Inches(3.3), Inches(1.2),
                 desc, font_size=15, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 15: FUTURE ROADMAP
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_PURPLE)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(6), Inches(0.7),
             "FUTURE ROADMAP", font_size=36, color=ACCENT_PURPLE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2.5), ACCENT_PURPLE)

roadmap = [
    ("NOW", "Hackathon MVP", [
        "8 AI agents (4 Red + 4 Blue)",
        "4 LLM providers supported",
        "Battle Royale & Cross-Exam",
        "Real-time dashboard",
        "5 compliance frameworks",
    ], ACCENT_GREEN),
    ("NEXT", "Post-Hackathon", [
        "CI/CD pipeline integration",
        "GitHub Actions / GitLab CI",
        "Scheduled recurring scans",
        "Team collaboration features",
        "Custom rule definitions",
    ], ACCENT_CYAN),
    ("FUTURE", "Enterprise Scale", [
        "SaaS platform deployment",
        "Multi-repo organization scan",
        "Historical trend analytics",
        "SIEM/SOAR integration",
        "SOC analyst copilot mode",
    ], ACCENT_PURPLE),
    ("VISION", "The End Game", [
        "Every PR auto-scanned",
        "Zero-day prediction engine",
        "Threat intelligence feeds",
        "Autonomous incident response",
        "AI security mesh network",
    ], ACCENT_ORANGE),
]

for i, (phase, subtitle, items, color) in enumerate(roadmap):
    x = Inches(0.4 + i * 3.2)
    y = Inches(1.5)
    card = add_shape(slide, x, y, Inches(3.0), Inches(5.3), fill_color=BG_CARD, border_color=color, border_width=Pt(2))

    add_text_box(slide, x + Inches(0.2), y + Inches(0.2), Inches(2.6), Inches(0.5),
                 phase, font_size=26, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.2), y + Inches(0.7), Inches(2.6), Inches(0.4),
                 subtitle, font_size=14, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)
    add_accent_line(slide, x + Inches(0.4), y + Inches(1.15), Inches(2.2), color)

    lines = [(f"  {item}", 13, LIGHT_GRAY) for item in items]
    add_multi_text(slide, x + Inches(0.3), y + Inches(1.3), Inches(2.5), Inches(3.5), lines, default_size=13)

    # Arrow between phases
    if i < 3:
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + Inches(3.02), y + Inches(2.3), Inches(0.2), Inches(0.4))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = DIM_GRAY
        arrow.line.fill.background()
        arrow.shadow.inherit = False


# ═══════════════════════════════════════════════════════════════════════════
# SLIDE 16: THANK YOU / CALL TO ACTION
# ═══════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=ACCENT_CYAN)

# Shield icon
add_shape(slide, Inches(5.9), Inches(0.8), Inches(1.5), Inches(1.5),
          fill_color=None, border_color=ACCENT_CYAN, border_width=Pt(3),
          shape_type=MSO_SHAPE.DIAMOND)

# Title
add_text_box(slide, Inches(1), Inches(2.5), Inches(11.3), Inches(1.0),
             "ZERO DAY FORTRESS", font_size=48, color=WHITE, bold=True,
             alignment=PP_ALIGN.CENTER)
add_accent_line(slide, Inches(4.5), Inches(3.5), Inches(4.3), ACCENT_CYAN)

# Subtitle
add_text_box(slide, Inches(2), Inches(3.8), Inches(9.3), Inches(0.6),
             "The future of autonomous cybersecurity is here.",
             font_size=22, color=ACCENT_CYAN, alignment=PP_ALIGN.CENTER)

# Stats row
stats = [
    ("8", "AI Agents", ACCENT_RED),
    ("4", "LLM Providers", ACCENT_PURPLE),
    ("12+", "Vuln Types", ACCENT_ORANGE),
    ("5", "Compliance\nFrameworks", ACCENT_GREEN),
    ("7", "Battle\nPhases", ACCENT_CYAN),
]

for i, (num, label, color) in enumerate(stats):
    x = Inches(1.0 + i * 2.3)
    y = Inches(4.6)
    card = add_shape(slide, x, y, Inches(2.0), Inches(1.4), fill_color=BG_CARD, border_color=color, border_width=Pt(1))
    add_text_box(slide, x, y + Inches(0.1), Inches(2.0), Inches(0.6),
                 num, font_size=32, color=color, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x, y + Inches(0.7), Inches(2.0), Inches(0.6),
                 label, font_size=12, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

# GitHub link
add_text_box(slide, Inches(2), Inches(6.3), Inches(9.3), Inches(0.5),
             "github.com/NavaNaveen/Zero-day-fortess",
             font_size=16, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)

# Thank you
add_text_box(slide, Inches(2), Inches(6.8), Inches(9.3), Inches(0.5),
             "Thank you! Questions?",
             font_size=20, color=ACCENT_CYAN, bold=True, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════
# SAVE
# ═══════════════════════════════════════════════════════════════════════════
output_path = "/Users/navaselvam/Desktop/zero-day-fortress/Zero_Day_Fortress_Hackathon_Presentation.pptx"
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
print(f"Total slides: {len(prs.slides)}")
